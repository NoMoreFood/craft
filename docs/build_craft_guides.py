"""Build the Word guides from the operating-mode diagrams and setup references.

Run from the repository root after docs/build_craft_diagrams.py.
Requires python-docx, pdf2image and Poppler; --poppler-path selects its binary directory.
Export the resulting DOCX files to PDF with Word and inspect the rendered pages.
"""

import argparse
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pdf2image import convert_from_path


ROOT = Path(__file__).resolve().parents[1]
REPO = "https://github.com/NoMoreFood/craft/blob/main/"
SETUP = [
    "source/README.md#mode-1-setup-privileged-certificate-enrollment",
    "source/README.md#mode-2-setup-privileged-credential-linking",
    "source/README.md#mode-3-setup-unprivileged-windows-exported-certificate",
    "source/README.md#mode-4-setup-service-tickets-through-constrained-delegation",
]
FLOW = [
    "docs/diagrams/02-enrollment-sequence.svg",
    "docs/diagrams/03-credential-linking-sequence.svg",
    "docs/diagrams/04-windows-exported-certificate.svg",
    "docs/diagrams/07-constrained-delegation.svg",
]


def hyperlink(paragraph, label, target):
    url = target if target.startswith("https://") else REPO + target
    relation = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "315DC0")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def links(doc, items):
    p = doc.add_paragraph()
    for index, (label, target) in enumerate(items):
        if index:
            p.add_run("  |  ")
        hyperlink(p, label, target)
    return p


def dimensions(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11 if landscape else 8.5)
    section.page_height = Inches(8.5 if landscape else 11)
    section.left_margin = section.right_margin = Inches(0.5 if landscape else 0.68)
    section.top_margin = section.bottom_margin = Inches(0.5 if landscape else 0.62)
    section.header_distance = section.footer_distance = Inches(0.25)


def new_document(path, title, subtitle):
    doc = Document(path)
    for child in list(doc._element.body):
        if child.tag != qn("w:sectPr"):
            doc._element.body.remove(child)
    for relation in list(doc.part.rels.values()):
        if relation.reltype.endswith(("/hyperlink", "/image")):
            doc.part.drop_rel(relation.rId)
    dimensions(doc.sections[0])
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption", "Small Note"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        props = style.element.find(qn("w:pPr"))
        if props is not None:
            for border in list(props.findall(qn("w:pBdr"))):
                props.remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (("Title", 25), ("Heading 1", 16), ("Heading 2", 12), ("Caption", 8.5)):
        doc.styles[name].font.size = Pt(size)
    doc.styles["Title"].paragraph_format.space_after = Pt(7)
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].paragraph_format.space_before = Pt(10)
        doc.styles[name].paragraph_format.space_after = Pt(5)
    code = doc.styles["Code Block"]
    code.font.name = "Consolas"
    code.font.size = Pt(9)
    code.font.color.rgb = RGBColor(0, 0, 0)
    code.paragraph_format.line_spacing = 1.08
    code.paragraph_format.space_after = Pt(8)
    doc.core_properties.title = title
    doc.core_properties.subject = "CRAFT Linux operating modes and PKINIT setup"
    doc.core_properties.author = "CRAFT"
    for section in doc.sections:
        for part in (section.header, section.footer):
            for child in list(part._element):
                part._element.remove(child)
        p = section.header.add_paragraph("CRAFT | " + title)
        p.style = doc.styles["Small Note"]
        p.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        p.runs[0].font.size = Pt(9)
        p = section.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        p._p.append(field)
    doc.add_paragraph(title, "Title")
    doc.add_paragraph(subtitle, "Subtitle")
    return doc


def paragraph(doc, text):
    return doc.add_paragraph(text)


def code(doc, text):
    p = doc.add_paragraph(text.strip(), "Code Block")
    p.paragraph_format.keep_together = True
    return p


def page(doc, title):
    doc.add_heading(title, 1).paragraph_format.page_break_before = True


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for column, width in zip(t.columns, widths, strict=True):
        column.width = Inches(width)
    for cell, label in zip(t.rows[0].cells, headers, strict=True):
        cell.text = label
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            if isinstance(value, tuple):
                hyperlink(cell.paragraphs[0], value[0], value[1])
            else:
                cell.text = value
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn("w:" + key), value)
        borders.append(border)
    t._tbl.tblPr.append(borders)
    for index, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                margin = OxmlElement("w:" + side)
                margin.set(qn("w:w"), "100")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "18334A" if index == 0 else "F3F6FA" if index % 2 else "FFFFFF")
            props.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(10)
                    if index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def figure_page(doc, mode, title, description, image):
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    dimensions(section, landscape=True)
    doc.add_heading(title, 1)
    links(doc, [("Complete setup instructions", SETUP[mode - 1]), ("Editable flow diagram", FLOW[mode - 1])])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    picture = BytesIO()
    image.save(picture, format="PNG")
    picture.seek(0)
    shape = p.add_run().add_picture(picture, width=Inches(9.5))
    shape._inline.docPr.set("descr", description)
    caption = doc.add_paragraph(description, "Caption")
    caption.paragraph_format.space_after = Pt(0)


def license_and_references(doc):
    doc.add_heading("References and license", 2)
    links(doc, [("Install and configure", "source/README.md"), ("Security and operations", "source/SECURITY.md"),
                ("Acceptance checks", "source/TESTING.md")])
    links(doc, [("PKINIT specification", "https://www.rfc-editor.org/rfc/rfc4556"),
                ("MIT PKINIT configuration", "https://web.mit.edu/kerberos/krb5-latest/doc/admin/pkinit.html")])
    p = doc.add_paragraph((
        'Source code, scripts, configuration examples and documentation are '
        'licensed under GNU GPL version 3 only (GPL-3.0-only). See '
    ), "Small Note")
    hyperlink(p, "LICENSE", "LICENSE")
    p.add_run(" for the terms and warranty disclaimer.")


def process_guide(images):
    path = ROOT / "docs/01-CRAFT-Process-Overview.docx"
    doc = new_document(path,
        "CRAFT Linux Operating Modes and Tickets",
        "Certificate Request Agent for Tickets  |  Process overview")
    paragraph(doc, (
        "CRAFT obtains Kerberos credentials from Linux without storing the user's AD password. "
        "Modes 1-3 obtain a user TGT through PKINIT. Mode 4 uses broker constrained delegation "
        "to obtain only approved service tickets. Domain membership is not required."
    ))
    doc.add_heading("Choose an operating mode", 1)
    table(doc, ["Mode", "Credential source and authority", "Setup"], [
        ["1  Privileged enrollment", (
            'Linux generates a user key and enrolls a certificate through CES and AD '
            'CS with an enrollment-agent signature. The CA restricts the agent, '
            'template and recipients.'
        ), ("Mode 1 setup", SETUP[0])],
        ["2  Privileged credential linking", (
            "Linux adds a temporary public key to the user's msDS-KeyCredentialLink "
            'and uses Key Trust PKINIT. AD attribute delegation defines the target '
            'scope.'
        ), ("Mode 2 setup", SETUP[1])],
        ["3  Unprivileged Windows export", (
            'Windows enrolls and exports a certificate and matching private key. '
            "Linux uses the securely transferred PEM pair under the caller's UID."
        ), ("Mode 3 setup", SETUP[2])],
        ["4  Service tickets only", (
            'A dedicated broker uses S4U2Self and S4U2Proxy. AD restricts the target services; '
            'the user receives no TGT or broker credentials.'
        ), ("Mode 4 setup", SETUP[3])],
    ], [1.65, 4.15, 1.34])
    doc.add_heading("How CRAFT selects the mode", 1)
    paragraph(doc, (
        'mechanism=constrained_delegation selects mode 4 and ignores home certificates. '
        'Otherwise a complete NSS-home ~/.config/craft/user.pem and user.key pair selects '
        'mode 3. Partial or invalid pairs fail. Only an absent pair selects enrollment '
        'or Key Trust from mechanism; failures never switch mechanisms.'
    ))
    paragraph(doc, (
        'Mode 3 uses ordinary executables. Modes 1, 2 and 4 require the setuid launcher, '
        'dedicated craft account, craft-users caller group and private /run/craft. '
        'Users run craft directly without sudo. Network work runs outside root; '
        'the caller publishes the cache. CRAFT does not perform an MFA challenge.'
    ))
    links(doc, [("Mode selection diagram", "docs/diagrams/01-system-overview.svg"),
                ("Build and install", "source/README.md#build")])
    figure_page(doc, 1, "Mode 1 privileged certificate enrollment", (
        'Linux resolves the user, enrolls a restricted certificate through CES '
        "and AD CS, then obtains the user's TGT through PKINIT. The generated "
        'user pair is temporary; CA records remain.'
    ), images[0])
    figure_page(doc, 2, "Mode 2 privileged credential linking", (
        'Linux uses a delegated directory write and Key Trust PKINIT against the '
        'same writable DC. Cleanup must succeed before cache publication. A '
        'residual key after an outage needs administrator removal.'
    ), images[1])
    figure_page(doc, 3, "Mode 3 unprivileged Windows exported certificate", (
        'Windows enrollment and export precede secure PEM transfer. Linux '
        "validates the supplied pair and obtains the user's TGT as the caller. "
        'The PEM pair remains and must be replaced before expiry.'
    ), images[2])
    figure_page(doc, 4, "Mode 4 service tickets through constrained delegation", (
        'A private broker TGT authorizes protocol transition and constrained delegation. '
        'Only validated tickets for every configured service are published. Home PEM files '
        'are ignored; there is no PKINIT or directory write.'
    ), images[3])
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    dimensions(section)
    doc.add_heading("Identity trust and credential storage", 1)
    paragraph(doc, (
        'All modes request <Linux username>@<configured realm> and verify the '
        'returned credential principals, flags, AES encryption and lifetimes. Linux names'
        ' must match the intended AD sAMAccountName. Enrollment and credential '
        'linking resolve and check the directory UPN. The supplied-certificate '
        'path performs no LDAP lookup; the KDC must map that certificate to the '
        'fixed caller principal. The certificate CN cannot select another '
        'account.'
    ))
    table(doc, ["Mode", "Client credential and authorization", "What persists"], [
        ["1  Enrollment", (
            'CA-issued user certificate; matching key, UPN, template, three client '
            'EKUs, short validity, CA chain and CRLs. CA recipient restrictions '
            'authorize issuance.'
        ), (
            'Private agent/keytab/transport credentials stay under /etc/craft. CA '
            'issuance and audit records remain. Generated user key and certificate '
            'are released.'
        )],
        ["2  Credential linking", (
            'Locally issued certificate carrying a fresh RSA public key (default '
            '3072 bits). The KDC'
            ' matches the linked key; a trusted user-certificate CA chain is '
            'unnecessary. KDC trust is still required.'
        ), (
            'Service keytab and delegated AD rights remain. CRAFT removes its exact '
            'directory entry and releases the generated identity. Failed cleanup can '
            'leave a standing credential.'
        )],
        ["3  Supplied certificate", (
            'Windows-exported PEM pair; matching key, logon usage, UPN SAN, validity,'
            ' CA trust and CRLs. KDC account mapping authorizes authentication.'
        ), (
            "user.pem and user.key stay in the caller's home. CRAFT does not renew, "
            'overwrite or delete them. Replace them before expiry.'
        )],
        ["4  Service tickets", (
            'Broker AES keytab, protocol transition and explicit AD target SPNs. '
            'No certificate authentication or directory write.'
        ), (
            'Broker keytab stays protected under /etc/craft. Only final user service '
            'tickets and their session keys persist in the user cache.'
        )],
    ], [1.1, 3.1, 2.94])
    paragraph(doc, (
        'Every mode atomically publishes a caller-owned mode-0600 .krb5cc_craft '
        'in the NSS home. It contains a user TGT in modes 1-3, or only selected service tickets in mode 4,'
        " with session keys but no certificate, private key or broker evidence. Modes 1-3 use it to "
        'request service tickets. Success replaces the cache; acquisition failure'
        ' preserves it, and expiration does not delete it automatically.'
    ))
    paragraph(doc, (
        'Generated credentials and sealed PKINIT memory files are released after '
        'the operation. Memory can reach swap or privileged host capture. '
        'Deleting keys, removing a linked credential, ending a process or '
        'removing a cache does not revoke an already-issued ticket. Host and '
        'service-account compromise expose the authority granted to that '
        'installation.'
    ))
    links(doc,
        [("Enrollment-agent restrictions",
        "docs/03-Enrollment-Agent-Restrictions.md"),
        ("Key Trust delegation",
        "docs/04-Key-Credential-Link-Delegation.md"),
        ("Broker delegation", "docs/05-Constrained-Delegation.md")])
    page(doc, "Ticket lifetimes and everyday use")
    paragraph(doc, (
        'Modes 1-3 request a ten-hour initial TGT and a seven-day renewal window '
        'with AES256/AES128. The KDC controls the grant; certificate/key lifetime'
        ' can constrain initial validity. A renewal window is the period during '
        'which an eligible TGT can be renewed. It does not make one ticket valid '
        'for seven days.'
    ))
    paragraph(doc, (
        'require_full_tgt_lifetime=yes rejects shortened initial or renewal '
        'grants without replacing the cache. Administrators can set no to accept '
        'shorter grants with a warning. Returned ticket times are never '
        'rewritten. Enrollment certificate caps default to ten hours total, '
        "including backdating; they do not cap mode 3's supplied certificate or "
        "mode 2's local identity."
    ))
    doc.add_heading("Select the user cache", 2)
    code(doc, 'cache=$(/usr/local/bin/craft) || exit "$?"\nexport KRB5CCNAME="$cache"\nklist -ef -c "$KRB5CCNAME"')
    paragraph(doc, (
        'Kerberos-aware applications launched from this shell use the selected '
        'cache for SMB, LDAP and Kerberos-enabled web services. The service still'
        " enforces the user's permissions. CRAFT cannot change its parent shell's"
        ' environment; the optional with-craft wrapper runs a command with its '
        'cache selected.'
    ))
    doc.add_heading("Maintain a long running job", 2)
    code(doc, (
        'job_pid=$$\ncache=$(/usr/local/bin/craft-maintain --watch-pid "$job_pid")'
        ' || exit "$?"\nexport KRB5CCNAME="$cache"'
    ))
    paragraph(doc, (
        'craft-maintain returns when credentials and detached maintenance are ready. '
        'Modes 1-3 renew from the cached TGT and authenticate afresh near the renewal deadline. '
        'Mode 4 refreshes the complete service batch through its broker before the earliest expiry. '
        'Fresh authentication always uses the configured mode selection.'
    ))
    paragraph(doc, (
        'Install the maintainer without setuid/setgid. Protect system Kerberos configuration, '
        'and watch a shell or controller that lasts for the job. The scheduler must retain '
        'background processes. Modes 1, 2 and 4 need setuid access for fresh authentication; '
        'home refresh works under no_new_privs. --max-duration 21d adds a cutoff. '
        'Concurrent jobs for one UID share the cache and acquisition cooldown.'
    ))
    code(doc, '/usr/local/bin/craft-maintain --status --watch-pid "$job_pid"\nklist -ef -c "$KRB5CCNAME"')
    paragraph(doc, (
        'Failures preserve credentials and retry; later failures do not stop the '
        'job. Status and AUTHPRIV logs report failures. Maintenance stops with '
        'the watched PID and leaves the cache. Applications must reload refreshed'
        ' credentials. Keep CRLs current and replace the supplied certificate '
        'before expiry; TGT maintenance does not renew it.'
    ))
    links(doc,
        [("Maintenance setup and status",
        "source/README.md#long-running-jobs"),
        ("Credential timing diagram",
        "docs/diagrams/05-credential-timing.svg")])
    page(doc, "Operation and authorization")
    paragraph(doc, (
        'Mode 4 uses a dedicated broker with protocol transition and account-based constrained '
        'delegation to an exact list of target SPNs. AD limits target services, not the users '
        'the broker can impersonate. Protect the keytab, restrict local craft-users membership '
        'and control Linux UID/name assignment. Exclude sensitive users through domain '
        'delegation policy and test denial against the deployed KDC.'
    ))
    paragraph(doc, (
        'Services sharing an AD account and key are not isolated by their SPN labels. Use '
        'separate target service identities where isolation matters. Applications must request '
        'the configured service names exactly; aliases or additional services need explicit '
        'configuration and AD authorization. A service-only cache cannot acquire arbitrary tickets.'
    ))
    paragraph(doc, (
        'All modes require a working host logging socket at /dev/log. Administrators '
        'maintain CRLs, review authorization scope and protect persistent credentials. '
        'Setting enabled=no stops fresh acquisition; stop active maintainers separately '
        'to end renewal. Issued tickets and service sessions remain subject to domain policy.'
    ))
    paragraph(doc, (
        'Mode 4 refreshes all service tickets through the broker before the earliest expiry. '
        'It has no cache-only renewal and needs the broker and KDC for each refresh. '
        'Status identifies ticket_kind and ticket_expires. Removing broker delegation blocks fresh '
        'acquisition but does not revoke already-issued tickets.'
    ))
    links(doc, [("Audit and troubleshooting", "source/README.md#audit-records-and-diagnostics"),
                ("Operational disable", "source/SECURITY.md#immediate-disable")])
    license_and_references(doc)
    doc.save(path)
    print(f"Created {path}")


def configuration_guide():
    path = ROOT / "docs/02-CRAFT-Configuration-and-Validation.docx"
    doc = new_document(path, "CRAFT Configuration and Validation", (
        'Certificate Request Agent for Tickets  |  Administrator reference'
    ))
    paragraph(doc, (
        'Configure one of four Linux operating modes: privileged certificate '
        'enrollment through CES and AD CS, privileged credential linking through '
        'msDS-KeyCredentialLink, or unprivileged PKINIT with a certificate '
        'exported from Windows, or service-only constrained delegation with a broker keytab. This guide gives '
        'mode-specific configuration and acceptance checks. The linked installation '
        'guide contains complete Linux commands and the configuration reference.'
    ))
    table(doc, ["Requirement", "1  Enrollment", "2  Credential linking", "3  Windows export"], [
        ["Linux acquisition",
            "Setuid launcher; craft service account",
            "Setuid launcher; craft service account",
            "Caller UID; mode-0755 executables"],
        ["User credential",
            "CA-enrolled short-lived certificate",
            "Temporary linked public key and local identity",
            "Transferred certificate and private-key PEM pair"],
        ["Service keytab / LDAP", "Required", "Required", "Unused"],
        ["CES / agent / user template", "Required", "Unused", "Direct Windows enrollment template; no Linux CES/agent"],
        ["Authorization",
            "CA agent/template/recipient restrictions",
            "Scoped AD attribute delegation",
            "User enrollment and KDC certificate mapping"],
        ["KDC certificate trust / CRLs", "Required", "Required", "Required"],
        ["Setup",
            ("Mode 1 instructions",
            SETUP[0]),
            ("Mode 2 instructions",
            SETUP[1]),
            ("Mode 3 instructions",
            SETUP[2])],
    ], [1.72, 1.8, 1.8, 1.82])
    paragraph(doc, (
        'Mode 4 uses the same privileged Linux installation as modes 1 and 2, plus a dedicated '
        'broker AES keytab with protocol transition and approved target SPNs. It requires no '
        'certificates, CRLs, PKINIT, LDAP or CES. See the mode 4 setup page.'
    ))
    page(doc, "Shared Linux preparation")
    paragraph(doc, (
        'Build as an ordinary user from the repository root with GCC/libstdc++ 13.1 '
        'or equivalent C++20 support, CMake 3.16 or newer, OpenSSL 1.1.1 or newer, MIT Kerberos, libcurl, '
        'libxml2, OpenLDAP, SASL and MIT GSSAPI. Install MIT PKINIT for modes 1-3 and, for modes '
        '1 and 2, the SASL GSSAPI plugin. CES Negotiate needs libcurl SPNEGO. '
        'Use supported patched packages meeting '
        "the installation guide's requirements. A working /dev/log socket is required."
    ))
    code(doc, (
        'cmake -S source -B build -DCMAKE_BUILD_TYPE=Release\ncmake --build build '
        '-j2\nctest --test-dir build --output-on-failure'
    ))
    paragraph(doc, (
        'Install the production launcher, worker and optional maintainer manually; '
        'there is no CMake installation target. Keep '
        '/etc/craft/config and krb5.conf root-owned mode 0644 with protected '
        'parents. Pin the realm, KDC and expected KDC hostname; provide populated'
        ' kdc-trust.pem and current kdc-crls.pem for modes 1-3. Modes 1 and 3 also require '
        'ca-trust.pem and ca-crls.pem for the user certificate; enrollment also '
        'validates the agent. Empty placeholders are invalid. Refresh CRLs '
        'independently.'
    ))
    paragraph(doc, (
        'Modes 1, 2 and 4 add craft:craft, craft-users and /run/craft owned '
        'craft:craft mode 0700. Private service credentials are root:craft mode '
        '0640. Only the service account belongs to craft; ordinary users join '
        'craft-users. Keep enabled=no and the launcher non-setuid until the '
        'selected setup is ready, then enable acquisition and set launcher mode '
        '4750. Users invoke craft directly without sudo.'
    ))
    links(doc,
        [("Shared installation",
        "source/README.md#shared-linux-installation"),
        ("Privileged installation",
        "source/README.md#shared-privileged-installation"),
        ("Configuration reference",
        "source/README.md#configuration-reference")])
    page(doc, "Mode 1 privileged certificate enrollment setup")
    links(doc,
        [("Complete mode 1 setup",
        SETUP[0]),
        ("Enrollment flow diagram",
        FLOW[0]),
        ("CA authorization",
        "docs/03-Enrollment-Agent-Restrictions.md")])
    paragraph(doc, (
        'Complete the shared Linux and privileged installation. Provision a '
        'directory/submission account with AES keys and install submitter.keytab.'
        ' Install the enrollment-agent certificate and matching unencrypted key '
        'as agent.pem and agent.key. Install CES HTTPS trust as https-trust.pem. '
        'Protect these files as root:craft mode 0640.'
    ))
    paragraph(doc, (
        'Configure AD CS and CES for initial enrollment using a dedicated EOBO '
        'user template. Restrict the actual signing agent, permitted template and'
        ' approved recipient group at the CA; exclude privileged identities. '
        'Resolve Linux names to the intended AD sAMAccountName and directory UPN.'
        ' CA-built identity and SID mapping must satisfy current KDC policy. The '
        'CA helper performs partial lab provisioning and verifies narrow restrictions '
        'before publication/export; configure those restrictions first.'
    ))
    code(doc, """
enabled=no
domain=domain.local
realm=DOMAIN.LOCAL
mechanism=enrollment
netbios=DOMAIN
service_principal=svc-linux-enroll@DOMAIN.LOCAL
template_oid=1.3.6.1.4.1.311.21.8.999.1
ces_url=https://ces.domain.local/IssuingCA_CES_Kerberos/service.svc/CES
ces_auth=negotiate
key_bits=3072
certificate_cn={user}
tgt_seconds=36000
renew_seconds=604800
require_full_tgt_lifetime=yes
cert_remaining_max_seconds=36000
cert_total_max_seconds=36000
minimum_interval_seconds=60
""")
    paragraph(doc, (
        'Replace every placeholder. template_oid is the actual '
        'msPKI-Cert-Template-OID, not the template display name. The CES URL is '
        'pinned, not discovered through Windows enrollment policy. Configure the '
        'optional GC endpoint/base when needed. The helper signs DOMAIN\\user in '
        'the EOBO request and validates the returned certificate against the '
        'directory UPN.'
    ))
    table(doc, ["CES transport", "Additional requirements"], [
        ["negotiate", (
            'Service keytab authenticates HTTP as well as LDAP. Configure the HTTP '
            "SPN and required CES-to-CA constrained delegation. The worker's "
            'short-lived transport TGT is never returned to the user.'
        )],
        ["mtls", (
            'Use the actual certificate-authenticated CES endpoint and separate '
            'https-client.pem / https-client.key mapped to the submitting account. '
            'The directory keytab is still required.'
        )],
    ], [1.1, 6.04])
    paragraph(doc, (
        'CES HTTPS uses https-trust.pem for chain/hostname verification. Optional '
        'https-crls.pem enforces revocation for the whole server chain; without it, '
        'HTTPS CRL/OCSP checks are absent. Refresh all required CRLs independently.'
    ))
    paragraph(doc, (
        'Validate the template profile on the next page and prove CA refusal for '
        'out-of-scope recipients. Set enabled=yes and enable the setuid launcher '
        'only when ready. Test as an approved user with both home PEM files '
        'absent, so the supplied-certificate path does not hide the enrollment '
        'setup.'
    ))
    page(doc, "Mode 1 certificate and ticket policy")
    paragraph(doc, (
        'The CA controls the issued certificate. CSR extensions are requests, and'
        ' client-side validation cannot replace CA authorization. This profile '
        'applies to CRAFT enrollment; mode 3 accepts its separate '
        'supplied-certificate profile and mode 2 is authorized by its directory '
        'key.'
    ))
    table(doc, ["Property", "Required enrollment profile"], [
        ["Subject and identity", (
            'The UPN otherName SAN (1.3.6.1.4.1.311.20.2.3) must match the '
            'directory-resolved UPN. The CA supplies the account SID security '
            'extension for current AD certificate mapping. The CN is not the '
            'authentication identity.'
        )],
        ["Client EKUs", (
            'Client Authentication 1.3.6.1.5.5.7.3.2\nSmart Card Logon '
            '1.3.6.1.4.1.311.20.2.2\nPKINIT Client Authentication 1.3.6.1.5.2.3.4'
        )],
        ["Key and basic constraints", (
            'RSA defaults to 3072 bits; key_bits accepts 2048/3072/4096 and must meet '
            'the template minimum. Explicit CA:FALSE with no path length; '
            'digitalSignature required; keyCertSign and cRLSign forbidden. The CSR '
            'also requests keyEncipherment.'
        )],
        ["Template", (
            'Template-information extension must contain the configured template OID.'
            ' The signing agent must be authorized for that template and recipient.'
        )],
        ["Validity", (
            'Current validity with more than 60 seconds remaining. Default cap is ten'
            ' hours total notBefore-to-notAfter, including backdating, and ten hours '
            'remaining. The client refuses excessive validity rather than changing '
            'dates.'
        )],
        ["Trust and revocation", (
            'Matching private key, configured CA chain and current CRLs for the user '
            'certificate and enrollment agent. KDC certificate trust and revocation '
            'are checked separately.'
        )],
        ["Agent profile", (
            'Certificate Request Agent EKU 1.3.6.1.4.1.311.20.2.1, matching private '
            'key, trusted chain and current CRLs. Scope comes from CA restrictions, '
            'not that EKU.'
        )],
    ], [1.35, 5.79])
    doc.add_heading("Shared ticket policy", 2)
    paragraph(doc, (
        'tgt_seconds=36000 and renew_seconds=604800 request ten-hour initial '
        'validity and seven-day renewal in modes 1-3. Set renew_seconds to 0 to '
        'disable renewal; otherwise it must exceed tgt_seconds. The KDC decides the actual'
        ' grant, including certificate/key-lifetime constraints. Strict '
        'require_full_tgt_lifetime=yes rejects shortened grants and preserves the'
        ' cache; no accepts them with a warning. CRAFT never rewrites issued '
        'ticket times. AES256/AES128 are required for both session keys and outer'
        ' ticket encryption.'
    ))
    doc.add_heading("Generated common name", 2)
    paragraph(doc, (
        'certificate_cn={user} controls the mode 1 CSR CN and mode 2 '
        'local-certificate CN. Tokens are {user}, directory-resolved {upn} and '
        '{domain}; the result must be 1-64 printable ASCII characters. A CA-built'
        " subject can override the CSR. The setting does not alter mode 3's "
        'supplied certificate or the requested principal. Preserve UPN/SID '
        'binding and recipient restrictions when configuring CA subject policy.'
    ))
    links(doc,
        [("Windows enrollment preparation",
        "source/README.md#assumptions-and-windows-preparation"),
        ("Configuration example",
        "source/config/config.example")])
    page(doc, "Mode 2 privileged credential linking setup")
    links(doc,
        [("Complete mode 2 setup",
        SETUP[1]),
        ("Credential-linking diagram",
        FLOW[1]),
        ("Directory delegation",
        "docs/04-Key-Credential-Link-Delegation.md")])
    paragraph(doc, (
        'Complete shared Linux and privileged installation. Provision the service'
        ' principal with AES keys and install submitter.keytab as root:craft mode'
        ' 0640. Use a schema exposing msDS-KeyCredentialLink and a writable KDC '
        'supporting NGC/Key Trust, as provided by Windows Server 2016 or later. A'
        ' domain-functional-level label alone does not establish these '
        'capabilities. The KDC still needs a trusted PKINIT certificate and '
        'current CRLs.'
    ))
    paragraph(doc, (
        'Pin the same writable DC in krb5.conf and kt_dc_url so its KDC sees the '
        'new key immediately. The LDAP SASL/GSSAPI bind must provide integrity '
        'and encryption at SASL SSF 128 or higher. For Key Trust alone, remove the optional '
        'pkinit_pool reference to user-certificate ca-trust.pem if unnecessary, '
        'or point it at a root-controlled bundle of KDC intermediates; retain KDC'
        ' anchors and CRL checking.'
    ))
    code(doc, """
enabled=no
domain=domain.local
realm=DOMAIN.LOCAL
mechanism=key_trust
service_principal=svc-linux-enroll@DOMAIN.LOCAL
kt_dc_url=ldap://dc01.domain.local
key_bits=3072
tgt_seconds=36000
renew_seconds=604800
require_full_tgt_lifetime=yes
minimum_interval_seconds=60
""")
    paragraph(doc, (
        'Delegate ReadProperty/WriteProperty on only msDS-KeyCredentialLink over '
        'a reviewed OU of ordinary users or a single user. The service account '
        'can authenticate as every target in that scope. Keep privileged '
        'identities out, protect the OU membership/ACL administrators, and '
        'prevent the service account from rewriting its delegation. Preview the '
        'helper rule before applying it:'
    ))
    code(doc, (
        '.\\scripts\\Grant-CRAFTKeyCredentialLink.ps1 -ServiceAccount '
        'svc-linux-enroll `\n    -TargetOU "OU=CRAFT Users,DC=domain,DC=local" '
        '-WhatIf'
    ))
    paragraph(doc, (
        'CRAFT generates a fresh RSA key and locally issued certificate. '
        'The linked key authorizes PKINIT without user-certificate CA enrollment, '
        'CES, an agent or a template. MIT Kerberos handles freshness tokens; '
        'validate against the actual KDC policy. Review the temporary issuer trust scope '
        'in the linked security guide.'
    ))
    paragraph(doc, (
        "An independent process adds and removes CRAFT's exact value while "
        'preserving existing keys. Successful acquisition requires removal before'
        ' cache publication. Cleanup also runs after acquisition failure or '
        'worker termination and retries uncertain directory outcomes. A '
        'continuing directory outage or host failure can leave a usable key. '
        'Failed cleanup blocks publication and reports a CRITICAL AUTHPRIV event '
        'identifying the object and worker when cleanup can run; inspect and remove '
        'only confirmed residual values, retaining legitimate keys.'
    ))
    paragraph(doc, (
        'Review delegation and trust, set enabled=yes, and enable the setuid '
        'launcher. Test with both home PEM files absent. Verify the returned '
        'caller principal, attribute cleanup, existing-key preservation and '
        'directory refusal for out-of-scope accounts; repeat after revoking the '
        'delegation.'
    ))
    links(doc, [("Temporary issuer trust", "source/SECURITY.md#key-trust-certificate-handling")])
    page(doc, "Mode 3 unprivileged Windows exported certificate setup")
    links(doc,
        [("Complete mode 3 setup",
        SETUP[2]),
        ("Windows export flow diagram",
        FLOW[2]),
        ("Windows helper details",
        "source/README.md#request-a-home-certificate-on-windows")])
    paragraph(doc, (
        'Complete shared Linux installation with ordinary mode-0755 executables. '
        'No setuid bit, service account, caller group, service keytab, LDAP '
        'lookup, CES or /run/craft is needed for Linux acquisition. Host '
        'installation and public Kerberos/trust provisioning remain administrator'
        ' tasks.'
    ))
    paragraph(doc, (
        'On Windows, run the helper as the intended signed-in domain user without'
        " elevation. Choose the template's internal name or OID, with direct user"
        ' enrollment, exportable software keys and a KDC-accepted '
        'identity/profile. Windows PowerShell and its PKI module use AD '
        'enrollment policy to select an eligible CA. No OpenSSL or '
        'enrollment-agent credential is required.'
    ))
    code(doc, '.\\scripts\\Request-CRAFT-Certificate.cmd UserLogon "C:\\Users\\Alice\\CRAFT Certificate"')
    paragraph(doc, (
        'Securely transfer both exported user.pem and matching unencrypted PKCS#8'
        ' user.key to Linux after the helper returns successfully. Keep its '
        'private tracking metadata on Windows. As the intended Linux user, '
        'install the pair in the NSS-resolved home; HOME and XDG overrides do not'
        ' change that location:'
    ))
    code(doc, """
mkdir -p ~/.config/craft
chmod 0700 ~/.config/craft
install -m 0600 /path/to/transferred/user.pem ~/.config/craft/user.pem
install -m 0600 /path/to/transferred/user.key ~/.config/craft/user.key
""")
    paragraph(doc, (
        'Both files must be caller-owned regular files, nonempty, at most 1 MiB '
        'and have one hard link. Symlinks are refused. The key must be '
        'unencrypted with no group/other permissions; the certificate must not be'
        ' group/world-writable. The home, .config and craft directories must be '
        'caller-owned and not group/world-writable. Direct PFX loading, '
        'passphrase prompts, PKCS#11 and hardware-key backends are unsupported.'
    ))
    code(doc, (
        'enabled=no\ndomain=domain.local\nrealm=DOMAIN.LOCAL\ntgt_seconds=36000\n'
        'renew_seconds=604800\nrequire_full_tgt_lifetime=yes'
    ))
    paragraph(doc, (
        'Provide root-owned mode-0644 CA and KDC trust bundles with current CRLs.'
        ' The user certificate must match the key, be non-CA, and have digitalSignature '
        'without certificate/CRL signing, Smart Card Logon or PKINIT Client '
        'Authentication EKU, exactly one UPN SAN, and more than 60 seconds of current validity. basicConstraints '
        'may be absent; if present, require CA:FALSE without a path length. The KDC '
        'must map it to <Linux '
        'username>@<realm>. CRAFT does not query the directory UPN or pin the '
        'enrollment template/ten-hour certificate cap in this mode.'
    ))
    paragraph(doc, (
        'Set enabled=yes when ready. There is no mechanism=home setting: the '
        'complete PEM pair selects this path, including on privileged '
        'installations. Partial or invalid pairs fail without fallback. Repeat '
        'the Windows helper to check renewal and transfer replacements before '
        'expiry. Its default window is one calendar month; /RenewBeforeDays '
        'changes it, and /DeleteAfterExport optionally removes the Windows-store '
        'certificate. Neither the helper nor CRAFT schedules or transfers '
        'replacements automatically.'
    ))
    page(doc, "Mode 4 service tickets through constrained delegation setup")
    links(doc, [("Complete mode 4 setup", SETUP[3]), ("Broker AD policy", "docs/05-Constrained-Delegation.md"),
                ("Delegation flow", FLOW[3])])
    paragraph(doc, (
        'Use a dedicated ordinary broker account, unique SPN and AES keytab. Enable protocol '
        'transition and account-based constrained delegation to only the required services. '
        'TrustedForDelegation must remain false. Install submitter.keytab as root:craft mode 0640, '
        'and pin the realm/KDC and AES policy in root-controlled krb5.conf. No CA or LDAP setup is needed.'
    ))
    code(doc, """
enabled=no
domain=domain.local
realm=DOMAIN.LOCAL
mechanism=constrained_delegation
service_principal=CRAFT/broker.domain.local@DOMAIN.LOCAL
service_spns=cifs/files.domain.local@DOMAIN.LOCAL HTTP/portal.domain.local@DOMAIN.LOCAL
service_ticket_seconds=36000
minimum_interval_seconds=60
""")
    paragraph(doc, (
        'service_spns accepts 1-32 exact same-realm service/host@REALM names separated by spaces. '
        'Wildcards, krbtgt, the broker principal and cross-realm targets are refused. Every target '
        'must also be authorized in the broker account msDS-AllowedToDelegateTo list. One denied '
        'target fails the entire acquisition and preserves the previous cache. Home certificates are ignored.'
    ))
    paragraph(doc, (
        'service_ticket_seconds defaults to 36000 and accepts 60-604800 seconds. The KDC may '
        'shorten grants; CRAFT accepts shorter valid AES tickets. TGT lifetime and renewal settings '
        'do not apply. Final tickets may be forwardable/forwarded, but no TGT, broker credential, '
        'S4U2Self evidence or proxy metadata is published.'
    ))
    paragraph(doc, (
        'The AD policy limits services, not recipient users. Restrict local callers through craft-users '
        'and protect Linux name assignment. Protect sensitive AD users against delegation. Services '
        'sharing one AD account/key are not isolated by different SPN labels. Use separate target '
        'identities where service isolation matters.'
    ))
    paragraph(doc, (
        'Stop existing maintainers and remove the old TGT cache before switching to mode 4. '
        'After enabling acquisition, verify every cache record, real service authentication, '
        'unapproved-target and nondelegable-user denial, home-pair precedence and cache preservation. '
        'Existing credentials are not revoked by configuration changes.'
    ))
    paragraph(doc, (
        'craft-maintain refreshes the complete service batch before its earliest expiry, through the '
        'broker rather than TGT renewal. Keep the acquisition interval below the actual refresh interval. '
        'Status reports ticket_kind=service, earliest ticket_expires and renew_until=0. Test refresh, '
        'outages and watched-job exit with short service lifetimes.'
    ))
    page(doc, "Acceptance checks and ongoing operation")
    paragraph(doc, (
        'Run the offline checks, then validate the chosen installed mode against '
        'disposable identities in an isolated AD lab. Offline fixtures or a '
        'successful MIT KDC trial do not establish AD certificate mapping, CES '
        'interoperability, the directory delegation or the installed setuid '
        'boundary.'
    ))
    table(doc, ["Scope", "Required observations"], [
        ["Every mode", (
            'Run as the intended non-root user without sudo. Confirm Linux-name '
            'principal, caller-owned mode-0600 cache, '
            'mode-appropriate flags and lifetimes, AES session/ticket '
            'encryption and actual lifetimes. Test a real Kerberos-enabled service.'
        )],
        ["Selection and failure", (
            'Mode 4 ignores home certificates. In modes 1-3, a complete home pair takes priority. Invalid files must'
            ' preserve the cache without either privileged mechanism. With both '
            'absent, use only the configured mechanism. Test wrong mapping, stale '
            'required CRLs, key mismatch, KDC outage, RC4 and shortened strict-mode '
            'grants.'
        )],
        ["1  Enrollment", (
            'Inspect UPN, CA SID mapping, template, three EKUs, matching key, CN '
            'behavior, chain/CRLs and total validity. Prove CA denial for '
            'unauthorized agents/templates/recipients and correlate CA/CES audit '
            'identity.'
        )],
        ["2  Credential linking", (
            'Prove scoped attribute rights and denial outside the scope. Inspect the '
            'added/removed value and existing keys before and after success, failed '
            'PKINIT and interruption. Force cleanup failure, verify refusal to '
            'publish and remove the residual key. Test revocation of delegation and '
            'KDC freshness interoperability.'
        )],
        ["3  Windows export", (
            'Verify ordinary Linux execution without service credentials or runtime '
            'directory, safe Windows PEM export/transfer, KDC mapping, certificate '
            'replacement and inherited no_new_privs. Wrong-account and unsafe/partial'
            ' pairs must preserve the cache.'
        )],
    ], [1.3, 5.84])
    code(doc, 'cache=$(/usr/local/bin/craft) || exit "$?"\nexport KRB5CCNAME="$cache"\nklist -ef -c "$KRB5CCNAME"')
    doc.add_heading("Long running jobs", 2)
    code(doc, (
        'job_pid=$$\ncache=$(/usr/local/bin/craft-maintain --watch-pid "$job_pid")'
        ' || exit "$?"\nexport KRB5CCNAME="$cache"'
    ))
    paragraph(doc, (
        'The ordinary maintainer needs Linux with /proc and system /etc/krb5.conf '
        'matching the pinned realm/KDC and AES policy. It monitors the job through '
        'a pidfd when available, or /proc on older kernels such as RHEL 8. The '
        'scheduler must retain detached processes. Renewal uses the cache alone; '
        'fresh authentication applies the same configured mode selection. Keep CRLs, '
        'service credentials and supplied certificates current. Status/log '
        'failures preserve the cache and retry; later failures do not terminate '
        'the job.'
    ))
    paragraph(doc, (
        'Test renewal, fresh authentication near the absolute renewal deadline, '
        'outages, concurrent jobs and termination with the watched PID under the '
        'actual scheduler. Applications must reload refreshed credentials. Cache '
        'files remain after exit, and certificate expiry or directory-key removal'
        ' does not revoke already-issued tickets.'
    ))
    links(doc,
        [("Full acceptance cases",
        "source/TESTING.md"),
        ("Maintenance setup and status",
        "source/README.md#long-running-jobs")])
    page(doc, "Audit maintenance and operational disable")
    paragraph(doc, (
        'Read AUTHPRIV records from the host logging service with administrator access. '
        'The worker request= field matches the launcher PID. Key Trust records identify '
        'the object and worker; correlate them with directory change event 5136 and KDC '
        'event 4768. Configure the required directory SACL and CA audit categories. '
        'Do not place private keys, tickets or export passwords in logs.'
    ))
    code(doc, 'sudo journalctl -t craft -t craft-worker -t craft-maintain')
    paragraph(doc, (
        'Every mode requires /dev/log. Failure to submit an audit record blocks startup, '
        'credential release or renewal publication. The final launcher completion record '
        'follows cache replacement, so a logging failure there can be reported after the '
        'cache changed. Local submission does not establish durable retention or forwarding. '
        'Install the example host audit rules only after adjusting paths, architecture and '
        'rules for files actually present.'
    ))
    paragraph(doc, (
        'For failures, inspect the log identity/PID named in the caller notice. Check clocks, '
        'PKINIT/SASL plugins, principal and SPNs, mapping, current CRLs, certificate profiles '
        'and effective CA/directory authorization. Diagnose a rejection without weakening '
        'TLS verification, required CRLs or mapping policy.'
    ))
    table(doc, ["Operation", "Administrator action"], [
        ["Keep authentication available", (
            'Refresh CA/KDC and optional CES CRLs before next-update; rotate service '
            'credentials; replace and securely transfer supplied certificates before expiry.'
        )],
        ["Stop fresh acquisition", (
            'Set enabled=no or remove launcher execute access. Removing only setuid stops '
            'modes 1, 2 and 4; a supplied pair remains usable only in certificate modes.'
        )],
        ["Stop ticket renewal", (
            'Inspect job status and send SIGTERM to each recorded maintainer_pid. '
            'Acquisition disable alone does not stop cache-only renewal. Cache files remain.'
        )],
        ["Withdraw authentication authority", (
            'Revoke the CA agent/recipient authorization or AD delegation, rotate compromised '
            'credentials and remove confirmed residual Key Trust values. Review issued '
            'certificates, tickets and service sessions through incident response.'
        )],
    ], [1.65, 5.49])
    paragraph(doc, (
        'Assign owners for credential/CRL lifecycle, recipient/delegation review and logs. '
        'Repeat acceptance after policy or credential changes. Certificate/key removal, '
        'cache deletion and maintenance stop do not revoke already-issued tickets.'
    ))
    paragraph(doc, (
        'Mode 4 refreshes all service tickets through the broker before the earliest expiry. '
        'It has no cache-only renewal and needs the broker and KDC for each refresh. '
        'Status identifies ticket_kind and ticket_expires. Removing broker delegation blocks fresh '
        'acquisition but does not revoke already-issued tickets.'
    ))
    links(doc, [("Audit and troubleshooting", "source/README.md#audit-records-and-diagnostics"),
                ("Host audit rules", "source/config/craft.audit.rules"),
                ("Operational disable", "source/SECURITY.md#immediate-disable")])
    p = doc.add_paragraph((
        'Source code, scripts, configuration examples and documentation are '
        'licensed under GNU GPL version 3 only (GPL-3.0-only). See '
    ), "Small Note")
    hyperlink(p, "LICENSE", "LICENSE")
    p.add_run(" for the terms and warranty disclaimer.")
    doc.save(path)
    print(f"Created {path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--poppler-path", help="Directory containing pdftoppm and pdfinfo")
    args = parser.parse_args()
    images = convert_from_path(str(ROOT / "docs/CRAFT-Diagrams.pdf"), dpi=100,
                               first_page=2, last_page=4, poppler_path=args.poppler_path)
    images += convert_from_path(str(ROOT / "docs/CRAFT-Diagrams.pdf"), dpi=100,
                                first_page=7, last_page=7, poppler_path=args.poppler_path)
    process_guide(images)
    configuration_guide()
